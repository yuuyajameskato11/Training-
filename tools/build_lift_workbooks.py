#!/usr/bin/env python3
"""Rebuild the Lift Card workbooks around one MAXES sheet.

The originals repeat the maxes on every tab, so a new training max means
editing nine sheets by hand.  These rebuilds put six numbers (plus the
rounding increment) on a single MAXES sheet, name them, and point every
formula on every week at those names — change one cell, the whole block
re-prescribes itself.  Set the increment to 2.5 and type your maxes in kilos
and the card is in kilos.

    python3 tools/build_lift_workbooks.py            # -> programs/*.xlsx
"""
import argparse
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

INK = "0B0B0A"
BONE = "E9E5CF"
PAPER = "FBFAF5"
RULE = "D9D5C4"
MUTED = "6E6A5C"

BOOK = Font(name="Helvetica Neue", color="1A1A17")
TITLE = Font(name="Helvetica Neue", size=20, bold=True, color=BONE)
BRAND = Font(name="Helvetica Neue", size=9, bold=True, color=BONE)
DAY = Font(name="Helvetica Neue", size=10, bold=True, color=BONE)
MOVE = Font(name="Helvetica Neue", size=10, bold=True, color="1A1A17")
CUE = Font(name="Helvetica Neue", size=8.5, italic=True, color=MUTED)
SMALL = Font(name="Helvetica Neue", size=8.5, color=MUTED)
LOAD = Font(name="Helvetica Neue", size=11, bold=True, color="1A1A17")

FILL_INK = PatternFill("solid", fgColor=INK)
FILL_PAPER = PatternFill("solid", fgColor=PAPER)
FILL_INPUT = PatternFill("solid", fgColor="FFF6D6")
HAIRLINE = Border(bottom=Side(style="thin", color=RULE))

LIFTS = [
    ("clean", "Clean"),
    ("squat", "Squat"),
    ("bench", "Bench"),
    ("jerk", "Jerk"),
    ("snatch", "Snatch"),
    ("bw", "Bodyweight"),
]
NAMES = {"clean": "clean", "squat": "squat", "bench": "bench",
         "jerk": "jerk", "snatch": "snatch", "bw": "bodyweight"}


def load_data(path):
    src = open(path).read()
    return json.loads(src[src.index("{"): src.rindex(";")])


def formula(spec):
    """A prescription -> an Excel formula written against the named maxes."""
    if not spec:
        return None
    if "s" in spec:
        return spec["s"]
    if "n" in spec:
        return spec["n"]
    coef = repr(round(spec["c"], 6))
    if spec.get("sys"):
        return (f"=ROUND((squat+bodyweight)*{coef}/roundto,0)*roundto-bodyweight")
    return f"=ROUND({NAMES[spec['b']]}*{coef}/roundto,0)*roundto"


def maxes_sheet(ws, program):
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDE", (3, 24, 14, 34, 3)):
        ws.column_dimensions[col].width = width

    ws.merge_cells("B2:D2")
    ws["B2"] = "KNCT"
    ws["B2"].font = TITLE
    ws.merge_cells("B3:D3")
    ws["B3"] = f"Lift Cards — {program['name']} · {program['season']}"
    ws["B3"].font = BRAND
    for row in (2, 3):
        ws.row_dimensions[row].height = 28 if row == 2 else 18
        for col in "ABCDE":
            ws[f"{col}{row}"].fill = FILL_INK

    ws["B5"] = "YOUR TRAINING MAXES"
    ws["B5"].font = Font(name="Helvetica Neue", size=9, bold=True, color=MUTED)
    ws.merge_cells("B6:D6")
    ws["B6"] = ("Every weight on every week sheet is a percentage of these six "
                "numbers. Edit one and the whole program re-prescribes itself.")
    ws["B6"].font = CUE
    ws["B6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[6].height = 30

    first = 8
    for i, (key, label) in enumerate(LIFTS):
        row = first + i
        ws[f"B{row}"] = label
        ws[f"B{row}"].font = MOVE
        cell = ws[f"C{row}"]
        cell.value = program["defaults"].get(key, 0)
        cell.font = LOAD
        cell.fill = FILL_INPUT
        cell.number_format = "0.#"
        cell.alignment = Alignment(horizontal="center")
        for col in "BCD":
            ws[f"{col}{row}"].border = HAIRLINE
        ws.row_dimensions[row].height = 22

    row = first + len(LIFTS) + 1
    ws[f"B{row}"] = "Round to nearest"
    ws[f"B{row}"].font = MOVE
    ws[f"C{row}"] = 5
    ws[f"C{row}"].font = LOAD
    ws[f"C{row}"].fill = FILL_INPUT
    ws[f"C{row}"].alignment = Alignment(horizontal="center")
    ws[f"D{row}"] = "5 for pounds · 2.5 for kilos"
    ws[f"D{row}"].font = CUE
    for col in "BCD":
        ws[f"{col}{row}"].border = HAIRLINE

    tested = row + 2
    ws[f"B{tested}"] = "Tested maxes"
    ws[f"B{tested}"].font = Font(name="Helvetica Neue", size=9, bold=True, color=MUTED)
    ws.merge_cells(f"B{tested + 1}:D{tested + 1}")
    ws[f"B{tested + 1}"] = ("For reference — the card runs off the training maxes above, "
                            "which sit a little under a true one-rep max: "
                            + " · ".join(f"{label} {int(program['tested'][key])}"
                                         for key, label in LIFTS if key in program["tested"]))
    ws[f"B{tested + 1}"].font = CUE
    ws[f"B{tested + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[tested + 1].height = 30

    return {key: f"'{ws.title}'!$C${first + i}" for i, (key, _) in enumerate(LIFTS)}, \
           f"'{ws.title}'!$C${row}"


def week_sheet(ws, week):
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDE", (34, 5, 12, 12, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "portrait"
    ws.print_options.horizontalCentered = True

    ws["A1"] = week["label"] + (f" · {week['note']}" if week.get("note") else "")
    ws["A1"].font = Font(name="Helvetica Neue", size=13, bold=True, color=BONE)
    ws["E1"] = "KNCT"
    ws["E1"].font = BRAND
    ws["E1"].alignment = Alignment(horizontal="right")
    ws.row_dimensions[1].height = 24
    for col in "ABCDE":
        ws[f"{col}1"].fill = FILL_INK

    row = 3
    for day in week["days"]:
        ws[f"A{row}"] = day["name"].upper()
        ws[f"A{row}"].font = DAY
        for col, head in zip("BCDE", ("SET", "LOAD", "REPS", "DONE")):
            ws[f"{col}{row}"] = head
            ws[f"{col}{row}"].font = Font(name="Helvetica Neue", size=8, bold=True, color=BONE)
            ws[f"{col}{row}"].alignment = Alignment(horizontal="center")
        for col in "ABCDE":
            ws[f"{col}{row}"].fill = FILL_INK
        ws.row_dimensions[row].height = 18
        row += 1

        for block in day["blocks"]:
            ws[f"A{row}"] = block["label"]
            ws[f"A{row}"].font = MOVE
            ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="center")
            row += 1
            if block.get("notes"):
                ws.merge_cells(f"A{row}:E{row}")
                ws[f"A{row}"] = " · ".join(block["notes"])
                ws[f"A{row}"].font = CUE
                row += 1
            for i, s in enumerate(block["sets"], 1):
                ws[f"B{row}"] = i
                ws[f"B{row}"].font = SMALL
                ws[f"B{row}"].alignment = Alignment(horizontal="center")
                value = formula(s["w"])
                if value is not None:
                    ws[f"C{row}"] = value
                    ws[f"C{row}"].font = LOAD if str(value).startswith("=") else SMALL
                    ws[f"C{row}"].number_format = "0.#"
                    ws[f"C{row}"].alignment = Alignment(horizontal="center")
                ws[f"D{row}"] = s["r"]
                ws[f"D{row}"].font = SMALL
                ws[f"D{row}"].alignment = Alignment(horizontal="center")
                for col in "ABCDE":
                    ws[f"{col}{row}"].border = HAIRLINE
                row += 1
            row += 1
        row += 1
    return row


def build(program, out_dir):
    wb = Workbook()
    maxes = wb.active
    maxes.title = "MAXES"
    cells, roundto = maxes_sheet(maxes, program)

    for key, name in NAMES.items():
        wb.defined_names[name] = DefinedName(name, attr_text=cells[key])
    wb.defined_names["roundto"] = DefinedName("roundto", attr_text=roundto)

    for week in program["weeks"]:
        title = re.sub(r"[\\/*?:\[\]]", "", week["label"])[:31]
        week_sheet(wb.create_sheet(title), week)

    name = "KNCT-Lift-Cards-" + re.sub(r"[^A-Za-z0-9]+", "-", program["name"]).strip("-") + ".xlsx"
    path = os.path.join(out_dir, name)
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-d", "--data", default="js/liftcards-data.js")
    ap.add_argument("-o", "--out", default="programs")
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)
    for program in load_data(args.data)["programs"]:
        print("->", build(program, args.out))


if __name__ == "__main__":
    main()
