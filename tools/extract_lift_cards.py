#!/usr/bin/env python3
"""Turn the coach's Lift Card workbooks into the app's percentage-driven data.

Every weight in those workbooks is already a formula off the training-max row
(row 11: clean / squat / bench / jerk / snatch / bodyweight).  This reads the
formulas rather than the cached numbers, so the output is a *recipe*
(base lift x coefficient) that the app re-evaluates against your own maxes.

    python3 tools/extract_lift_cards.py <workbook.xlsx> [...] -o js/liftcards-data.js

Usage note: the source workbooks are not committed — point this at your own
copies to regenerate the data file.
"""
import argparse
import json
import re
import sys
from datetime import date

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl required:  pip install openpyxl")

MAX_ROW = 11          # training maxes drive every formula
TESTED_ROW = 10       # tested/goal maxes, shown for reference
COL_KEY = {"B": "clean", "C": "squat", "D": "bench", "E": "jerk", "F": "snatch", "G": "bw"}
DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")

# =ROUND($C$11*0.5*0.8/5,0/5)*5   -> squat x 0.40, rounded to nearest 5
SIMPLE = re.compile(r"^=ROUND\(\$([B-G])\$11((?:[*/][\d.]+)*)/5,0/5\)\*5$", re.I)
# =ROUND(($C$11+$G$11)*0.38/5,0/5)*5-$G$11  -> % of system weight, bar load only
SYSTEM = re.compile(r"^=ROUND\(\(\$C\$11\+\$G\$11\)\*([\d.]+)/5,0/5\)\*5-\$G\$11$", re.I)


def coefficient(chain):
    """Fold a '*0.5*0.8' / '/2*0.75' operator chain into one number."""
    value = 1.0
    for op, num in re.findall(r"([*/])([\d.]+)", chain):
        value = value * float(num) if op == "*" else value / float(num)
    return round(value, 6)


def parse_weight(raw):
    """A cell in the 'wt' column -> a computed prescription, a literal, or a cue."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return {"n": float(raw)}
    text = str(raw).strip()
    if not text:
        return None
    if text.startswith("="):
        m = SIMPLE.match(text)
        if m:
            return {"b": COL_KEY[m.group(1).upper()], "c": coefficient(m.group(2))}
        m = SYSTEM.match(text)
        if m:
            return {"b": "squat", "c": round(float(m.group(1)), 6), "sys": True}
        raise ValueError(f"unrecognised formula: {text}")
    return {"s": text}


def is_note(raw):
    """Cue lines (rest, pairing, tempo, coaching asides) vs. movement names."""
    text = str(raw)
    if not text.strip():
        return True
    if text[0].isspace() or text[0] in "(.":
        return True
    return text.strip()[0].islower()


def clean(raw):
    return re.sub(r"\s+", " ", str(raw)).strip() if raw is not None else ""


def read_day(ws, top, bottom, cols):
    """Rows top..bottom in one column trio -> blocks of {label, notes, sets}."""
    name_c, wt_c, rep_c = cols
    blocks, pending = [], []
    for row in range(top, bottom + 1):
        label_raw = ws[f"{name_c}{row}"].value
        weight = parse_weight(ws[f"{wt_c}{row}"].value)
        reps = clean(ws[f"{rep_c}{row}"].value)
        label = clean(label_raw)

        if label and not is_note(label_raw):
            blocks.append({"label": label, "notes": pending, "sets": []})
            pending = []
        elif label:
            if not weight and not reps:
                # a standalone cue: it introduces whatever comes next
                pending.append(label)
            elif blocks:
                if label not in blocks[-1]["notes"]:
                    blocks[-1]["notes"].append(label)

        if (weight or reps) and blocks:
            blocks[-1]["sets"].append({"w": weight, "r": reps})

    for b in blocks:
        if not b["notes"]:
            b.pop("notes")
        if re.search(r"\bmatrix\b", b["label"], re.I) and not b["sets"]:
            b["heading"] = True
    if pending and blocks:
        blocks[-1].setdefault("notes", []).extend(pending)
    return [b for b in blocks if b["sets"] or b.get("notes") or b.get("heading")]


def day_spans(ws, name_c):
    """Locate each 'Monday - warm up' header and the rows that belong to it."""
    heads = []
    for row in range(1, ws.max_row + 1):
        text = clean(ws[f"{name_c}{row}"].value)
        if any(text.lower().startswith(d.lower()) for d in DAYS) and "warm up" in text.lower():
            heads.append((row, text.split("-")[0].strip()))
    spans = []
    for i, (row, name) in enumerate(heads):
        end = heads[i + 1][0] - 1 if i + 1 < len(heads) else ws.max_row
        spans.append((name, row + 1, end))
    return spans


def read_maxes(ws, row):
    out = {}
    for col, key in COL_KEY.items():
        v = ws[f"{col}{row}"].value
        if isinstance(v, (int, float)):
            out[key] = float(v)
    return out


def week_label(ws):
    """Sheet E4 is the card's own name ('week 4'); title-case it, and keep the
    tab name too when it says something the header doesn't (e.g. Home Unload)."""
    label = clean(ws["E4"].value).title() or ws.title.title()
    tab = ws.title.title()
    squash = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())
    note = "" if squash(tab) in squash(label) or squash(label) in squash(tab) else tab
    return label, note


def read_sheet(ws):
    days = []
    for name_c, wt_c, rep_c in (("A", "B", "C"), ("E", "F", "G")):
        for name, top, bottom in day_spans(ws, name_c):
            blocks = read_day(ws, top, bottom, (name_c, wt_c, rep_c))
            if blocks:
                days.append({"name": name, "blocks": blocks, "_col": name_c, "_row": top})
    # left column first within a row band, then down the card
    days.sort(key=lambda d: (d["_row"], d["_col"]))
    order = {d: i for i, d in enumerate(DAYS)}
    days.sort(key=lambda d: order.get(d["name"], 99))
    for d in days:
        d.pop("_col"), d.pop("_row")
    label, note = week_label(ws)
    out = {
        "id": re.sub(r"[^a-z0-9]+", "-", ws.title.lower()).strip("-"),
        "label": label,
        "sheet": ws.title,
        "days": days,
    }
    if note:
        out["note"] = note
    return out


def read_workbook(path, program_id, name, blurb):
    wb = openpyxl.load_workbook(path, data_only=False)
    weeks = [read_sheet(ws) for ws in wb.worksheets]
    first = wb.worksheets[0]
    return {
        "id": program_id,
        "name": name,
        "blurb": blurb,
        "season": clean(first["A4"].value),
        "level": clean(first["F4"].value),
        "defaults": read_maxes(first, MAX_ROW),
        "tested": read_maxes(first, TESTED_ROW),
        "weeks": weeks,
    }


BOOKS = [
    ("adv", "Advanced", "Eight weeks of the advanced card — Olympic lifts, contrast work, matrices."),
    ("team", "Team · Phase 2", "Weeks 4-9 of the team block, closing with the home unload week."),
    ("dev", "Dev II · Weeks 1-4", "The development build: intro week plus the first three loading weeks."),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbooks", nargs="+")
    ap.add_argument("-o", "--out", default="js/liftcards-data.js")
    args = ap.parse_args()

    programs = []
    for path, (pid, name, blurb) in zip(args.workbooks, BOOKS):
        programs.append(read_workbook(path, pid, name, blurb))

    payload = {"generated": date.today().isoformat(), "programs": programs}
    body = json.dumps(payload, indent=1, ensure_ascii=False)
    with open(args.out, "w") as fh:
        fh.write("/* Generated by tools/extract_lift_cards.py — do not edit by hand. */\n")
        fh.write("window.LIFT_CARDS = ")
        fh.write(body)
        fh.write(";\n")

    for p in programs:
        sets = sum(len(b["sets"]) for w in p["weeks"] for d in w["days"] for b in d["blocks"])
        calc = sum(1 for w in p["weeks"] for d in w["days"] for b in d["blocks"]
                   for s in b["sets"] if s["w"] and "b" in s["w"])
        print(f"{p['id']:5} {len(p['weeks'])} weeks  "
              f"{sum(len(w['days']) for w in p['weeks'])} days  {sets} sets  {calc} computed")
    print("->", args.out)


if __name__ == "__main__":
    main()
